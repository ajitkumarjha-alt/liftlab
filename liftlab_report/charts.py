"""Native Excel charts, configured to be readable.

Every chart built here gets, without the caller having to ask:
  * visible tick labels on BOTH axes, with an explicit number format
  * horizontal major gridlines, vertical gridlines off
  * a y axis anchored at 0 unless the caller states a reason
  * a legend only when there is more than one series, placed clear of the plot
  * a deliberate size, so charts do not land on top of each other

openpyxl's defaults hide axes (delete=True) and omit tick marks, which is how
the v1 workbook ended up with axis titles but no numbers against them. The
axis configuration below is therefore explicit on every axis, every time.

Series data lives on the hidden DATA_CHARTS sheet: Excel charts reference cells,
so a chart with no backing cells is a chart with no numbers.
"""

from __future__ import annotations

from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.data_source import NumDataSource, NumRef
from openpyxl.chart.error_bar import ErrorBars
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.utils import get_column_letter

from . import stats

# Colours match the convention declared on READ THIS FIRST.
CLR_GOOD = "70AD47"        # at/below the sheet assumption
CLR_WARN = "FFC000"        # over the assumption, still within the cliff
CLR_BAD = "C00000"         # over the compliance cliff
CLR_NEUTRAL = "4472C4"
CLR_GREY = "A6A6A6"
CLR_LINE = "ED7D31"

# One standard chart footprint, so anchors can be computed instead of guessed.
CH_H, CH_W = 9.0, 18.0
CM_PER_ROW = 0.529         # a default-height Excel row

# Above this many categories the labels collide with their neighbours no matter
# how the chart is sized, so they are dropped and the value axis carries the
# numbers instead. An unreadable label is worse than no label.
MAX_LABELLED_CATS = 8


def rows_for(height_cm: float = CH_H, pad: int = 3) -> int:
    """Rows a chart of this height occupies, plus separation. Callers advance
    their row cursor by this so charts never overlap each other or a table."""
    return int(height_cm / CM_PER_ROW) + pad


class ChartData:
    """Sequential column blocks on a hidden sheet; hands back Reference ranges."""

    def __init__(self, wb):
        self.ws = wb.create_sheet("DATA_CHARTS")
        self.ws.sheet_state = "hidden"
        self.col = 1

    def block(self, cols: list[list]) -> tuple:
        """cols[i] = [header, v1, v2, ...]. Returns (ws, first_col, last_col,
        n_rows)."""
        c0 = self.col
        nrows = 0
        for j, colvals in enumerate(cols):
            for i, v in enumerate(colvals):
                self.ws.cell(row=i + 1, column=c0 + j, value=v)
            nrows = max(nrows, len(colvals))
        self.col = c0 + len(cols) + 1
        return self.ws, c0, c0 + len(cols) - 1, nrows

    def ref(self, first_col, last_col, first_row, last_row):
        return Reference(self.ws, min_col=first_col, max_col=last_col,
                         min_row=first_row, max_row=last_row)

    def abs_range(self, col, first_row, last_row) -> str:
        """An absolute A1 range on the data sheet — error bars need a string."""
        letter = get_column_letter(col)
        return f"'{self.ws.title}'!${letter}${first_row}:${letter}${last_row}"


# ── axis configuration ───────────────────────────────────────────────────────

def _axis(ch, x_title, y_title, x_fmt, y_fmt, y_min=0.0, y_max=None,
          gridlines=True):
    for ax, ttl, fmt in ((ch.x_axis, x_title, x_fmt), (ch.y_axis, y_title, y_fmt)):
        ax.title = ttl
        ax.delete = False                  # openpyxl hides axes by default
        ax.majorTickMark = "out"
        ax.minorTickMark = "none"
        ax.tickLblPos = "nextTo"
        if fmt:
            ax.numFmt = fmt
    ch.y_axis.majorGridlines = ChartLines() if gridlines else None
    ch.x_axis.majorGridlines = None        # vertical gridlines off
    if y_min is not None:
        ch.y_axis.scaling.min = y_min
    if y_max is not None:
        ch.y_axis.scaling.max = y_max
    return ch


def _data_labels(ch, n_cats: int, num_fmt: str = "#,##0", pos: str | None = "inEnd"):
    """Attach value-only data labels, or none at all.

    openpyxl OMITS unset DataLabelList attributes from the XML, and Excel then
    applies its OWN defaults — which print "series name, category name, value"
    against every point. Setting showVal=True alone is therefore not enough:
    every other flag must be written explicitly as False. That omission is what
    made the v2 charts unreadable, and on a stacked histogram it also printed a
    label for each of the two zero-height segments in every bin.

    Labels are attached ONLY where they can be read. Past MAX_LABELLED_CATS
    categories they are dropped and the value axis carries the numbers.
    """
    if n_cats > MAX_LABELLED_CATS:
        ch.dataLabels = None
        return ch
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showVal = True
    # Explicit False on every one of these — never left None.
    ch.dataLabels.showSerName = False
    ch.dataLabels.showCatName = False
    ch.dataLabels.showLegendKey = False
    ch.dataLabels.showPercent = False
    ch.dataLabels.showBubbleSize = False
    ch.dataLabels.numFmt = num_fmt
    if pos:
        ch.dataLabels.dLblPos = pos
    return ch


def _legend(ch, n_series: int, position="b"):
    """A legend only earns its space when there is more than one series, and it
    goes below or to the right so it never covers the plot."""
    if n_series > 1:
        ch.legend.position = position
        ch.legend.overlay = False
    else:
        ch.legend = None
    return ch


def _size(ch, h=CH_H, w=CH_W):
    ch.height, ch.width = h, w
    return ch


def _colour(series, rgb, line_only=False):
    gp = GraphicalProperties()
    if line_only:
        gp.line = LineProperties(solidFill=rgb, w=20000)
    else:
        gp.solidFill = rgb
        gp.line = LineProperties(solidFill=rgb)
    series.graphicalProperties = gp
    return series


# ── close-travel histogram, with the thresholds readable off the chart ───────

def close_histogram(cd: ChartData, title: str, values: list[float],
                    edges: list[float], assumption_s: float, cliff_s: float):
    """Histogram split into three SERIES by where each bin sits relative to the
    sheet assumption and the compliance cliff, so the legend names the edges
    instead of leaving the reader to count bars. Each bin carries a value in
    exactly one series, so the bars do not stack."""
    labels = stats.hist_labels(edges)
    counts = stats.hist(values, edges)
    lower = [0.0] + list(edges)            # lower bound of each bin

    at_or_below, over_assumption, over_cliff = [], [], []
    for i, _lab in enumerate(labels):
        lo = lower[i]
        n = counts[i]
        if lo >= cliff_s:
            at_or_below.append(None)
            over_assumption.append(None)
            over_cliff.append(n)
        elif lo >= assumption_s:
            at_or_below.append(None)
            over_assumption.append(n)
            over_cliff.append(None)
        else:
            at_or_below.append(n)
            over_assumption.append(None)
            over_cliff.append(None)

    wsd, c0, c1, nr = cd.block([
        ["close-travel band (s)"] + labels,
        [f"at or below the {assumption_s:.2f}s sheet assumption"] + at_or_below,
        [f"over {assumption_s:.2f}s, within the {cliff_s:.2f}s cliff"] + over_assumption,
        [f"OVER the {cliff_s:.2f}s compliance cliff"] + over_cliff])

    ch = BarChart()
    ch.type = "col"
    ch.grouping = "stacked"                # one value per bin, so bars align
    ch.overlap = 100
    ch.gapWidth = 40                       # wide bars, readable at this bin count
    ch.title = title
    ch.add_data(cd.ref(c0 + 1, c1, 1, nr), titles_from_data=True)
    ch.set_categories(cd.ref(c0, c0, 2, nr))
    for s, rgb in zip(ch.series, (CLR_GOOD, CLR_WARN, CLR_BAD)):
        _colour(s, rgb)
    _axis(ch, "close travel (s)", "number of closes", None, "#,##0")
    # 12 bins: past the labelling limit, so the value axis carries the numbers.
    _data_labels(ch, len(labels))
    _legend(ch, 3)
    return _size(ch, h=10.0, w=15.0)


# ── stopping rule ────────────────────────────────────────────────────────────

def stopping_rule(cd: ChartData, title: str, values: list[float], cliff_s: float):
    """Running mean close travel with its 95% band, against the cliff. The
    study can stop for this pool when the band clears the line and stays clear."""
    run_mean, lo_b, hi_b, cliff_line, idx = [], [], [], [], []
    s = s2 = 0.0
    for i, v in enumerate(values, start=1):
        s += v
        s2 += v * v
        m = s / i
        sd = (max(0.0, s2 - i * m * m) / (i - 1)) ** 0.5 if i > 1 else None
        half = (stats.Z95 * sd / (i ** 0.5)) if sd else None
        run_mean.append(round(m, 3))
        lo_b.append(round(m - half, 3) if half is not None else None)
        hi_b.append(round(m + half, 3) if half is not None else None)
        cliff_line.append(cliff_s)
        idx.append(i)
    wsd, c0, c1, nr = cd.block([
        ["closes measured (in order)"] + idx,
        ["running mean close travel"] + run_mean,
        ["lower end of the 95% range"] + lo_b,
        ["upper end of the 95% range"] + hi_b,
        [f"compliance cliff {cliff_s:.2f}s"] + cliff_line])
    ch = LineChart()
    ch.title = title
    ch.add_data(cd.ref(c0 + 1, c1, 1, nr), titles_from_data=True)
    ch.set_categories(cd.ref(c0, c0, 2, nr))
    for s_, rgb, dashed in ((ch.series[0], CLR_NEUTRAL, False),
                            (ch.series[1], CLR_GREY, True),
                            (ch.series[2], CLR_GREY, True),
                            (ch.series[3], CLR_BAD, False)):
        _colour(s_, rgb, line_only=True)
        s_.marker = Marker(symbol="none")
        s_.smooth = False
    ch.series[1].graphicalProperties.line.dashStyle = "dash"
    ch.series[2].graphicalProperties.line.dashStyle = "dash"
    _axis(ch, "closes measured (in order)", "close travel (s)", "#,##0", "0.00")
    _legend(ch, 4)
    return _size(ch)


# ── fleet comparison — the single most useful chart for a design reviewer ────

def fleet_comparison(cd: ChartData, rows: list[dict], cliff_s: float):
    """One bar per lift: median close travel with a 95% whisker, and the
    compliance cliff drawn straight across.

    rows: [{'label', 'median', 'err_lo', 'err_hi', 'n', 'over'}] — `over` marks
    a lift whose median sits above the cliff, which is coloured red."""
    if not rows:
        return None
    labels = [f"{r['label']} (n={r['n']})" for r in rows]
    meds = [round(r["median"], 3) for r in rows]
    lo = [round(r["err_lo"], 3) for r in rows]
    hi = [round(r["err_hi"], 3) for r in rows]
    cliff = [cliff_s] * len(rows)
    wsd, c0, c1, nr = cd.block([
        ["lift"] + labels,
        ["typical (median) close travel"] + meds,
        ["ci_minus"] + lo,
        ["ci_plus"] + hi,
        [f"compliance cliff {cliff_s:.2f}s"] + cliff])

    bar = BarChart()
    bar.type = "col"
    bar.title = "Typical door-close travel by lift, against the compliance cliff"
    bar.add_data(cd.ref(c0 + 1, c0 + 1, 1, nr), titles_from_data=True)
    bar.set_categories(cd.ref(c0, c0, 2, nr))
    series = bar.series[0]
    series.errBars = ErrorBars(
        errDir="y", errValType="cust", noEndCap=False,
        plus=NumDataSource(NumRef(f=cd.abs_range(c0 + 3, 2, nr))),
        minus=NumDataSource(NumRef(f=cd.abs_range(c0 + 2, 2, nr))))
    _colour(series, CLR_NEUTRAL)
    # Per-bar colouring: a lift over the cliff must read as over the cliff.
    for i, r in enumerate(rows):
        dp = DataPoint(idx=i)
        dp.graphicalProperties = GraphicalProperties(
            solidFill=CLR_BAD if r["over"] else CLR_GOOD)
        series.data_points.append(dp)
    bar.gapWidth = 40
    _data_labels(bar, len(rows), num_fmt="0.00", pos="inEnd")
    _axis(bar, "lift", "close travel (s)", None, "0.00")

    line = LineChart()
    line.add_data(cd.ref(c1, c1, 1, nr), titles_from_data=True)
    _colour(line.series[0], CLR_BAD, line_only=True)
    line.series[0].marker = Marker(symbol="none")
    line.series[0].smooth = False
    bar += line
    _legend(bar, 2)
    return _size(bar, h=10.0, w=20.0)


# ── generic bar / line ───────────────────────────────────────────────────────

def bar_chart(cd: ChartData, title, cats_ref, data_ref, n_series, n_cats,
              x_title, y_title, y_fmt="#,##0", colour=CLR_NEUTRAL,
              h=CH_H, w=CH_W, grouping=None):
    ch = BarChart()
    ch.type = "col"
    if grouping:
        ch.grouping = grouping
    ch.gapWidth = 40
    ch.title = title
    ch.add_data(data_ref, titles_from_data=True)
    ch.set_categories(cats_ref)
    if n_series == 1:
        _colour(ch.series[0], colour)
    _axis(ch, x_title, y_title, None, y_fmt)
    # Labels only survive when BOTH the category count and the series count
    # leave room: k series over n categories means k*n label boxes.
    _data_labels(ch, n_cats * max(1, n_series), num_fmt=y_fmt)
    _legend(ch, n_series)
    return _size(ch, h=h, w=w)


def line_chart(cd: ChartData, title, cats_ref, data_ref, n_series,
               x_title, y_title, y_fmt="#,##0", y_max=None, y_min=0.0):
    ch = LineChart()
    ch.title = title
    ch.add_data(data_ref, titles_from_data=True)
    ch.set_categories(cats_ref)
    for s in ch.series:
        s.marker = Marker(symbol="circle", size=5)
        s.smooth = False
    ch.dataLabels = None            # explicit: a 24-point line cannot carry them
    _axis(ch, x_title, y_title, None, y_fmt, y_min=y_min, y_max=y_max)
    _legend(ch, n_series, position="r")
    return _size(ch)


def demand_by_hour_grouped(cd: ChartData, hours: list[str], by_lift: dict,
                           title: str):
    """Grouped columns, hour x lift. A lift that was DARK in an hour carries
    None, not 0 — an unobserved hour must not draw a zero-height bar that reads
    as 'this lift carried nobody'."""
    cols = [["hour (IST)"] + hours]
    for label, vals in by_lift.items():
        cols.append([label] + [None if v is None else round(v, 2) for v in vals])
    wsd, c0, c1, nr = cd.block(cols)
    ch = BarChart()
    ch.type = "col"
    ch.grouping = "clustered"
    ch.gapWidth = 60
    ch.overlap = -10
    ch.title = title
    ch.add_data(cd.ref(c0 + 1, c1, 1, nr), titles_from_data=True)
    ch.set_categories(cd.ref(c0, c0, 2, nr))
    ch.dataLabels = None               # 24 hours x n lifts — never labellable
    _axis(ch, "hour of the day (IST)", "mean boardings per hour", None, "0.0")
    _legend(ch, len(by_lift), position="b")
    return _size(ch, h=10.0, w=24.0)


def fleet_demand_line(cd: ChartData, hours: list[str], fleet: list,
                      peak_hour: int | None, title: str):
    """Fleet total by hour, with the peak hour marked by a second series that
    is non-null only at the peak — so the mark cannot drift from the data."""
    peak_series = [None] * len(hours)
    if peak_hour is not None and 0 <= peak_hour < len(hours):
        peak_series[peak_hour] = fleet[peak_hour]
    cols = [["hour (IST)"] + hours,
            ["fleet total (sum of per-lift means)"]
            + [None if v is None else round(v, 2) for v in fleet],
            ["busiest hour"]
            + [None if v is None else round(v, 2) for v in peak_series]]
    wsd, c0, c1, nr = cd.block(cols)
    ch = LineChart()
    ch.title = title
    ch.add_data(cd.ref(c0 + 1, c1, 1, nr), titles_from_data=True)
    ch.set_categories(cd.ref(c0, c0, 2, nr))
    _colour(ch.series[0], CLR_NEUTRAL, line_only=True)
    ch.series[0].marker = Marker(symbol="circle", size=5)
    ch.series[0].smooth = False
    # the peak marker: a big dot, no connecting line
    ch.series[1].marker = Marker(symbol="diamond", size=11)
    gp = GraphicalProperties()
    gp.line = LineProperties(noFill=True)
    ch.series[1].graphicalProperties = gp
    ch.series[1].smooth = False
    ch.dataLabels = None
    _axis(ch, "hour of the day (IST)", "mean boardings per hour", None, "0.0")
    _legend(ch, 2, position="b")
    return _size(ch, h=9.0, w=22.0)


def coverage_timeline(cd: ChartData, days: list[str], by_cam: dict,
                      gap_share: list[float]):
    """Daily coverage per channel as lines, with the DECLARED-outage share of
    each day as grey bars behind them — so a dip caused by a known outage reads
    differently from a dip nobody has explained."""
    cols = [["day"] + days,
            ["share of the day inside a declared outage"] + gap_share]
    cams = sorted(by_cam)
    for cam in cams:
        # coverage is stored 0-100; charts want fractions to use a % format
        cols.append([f"{cam} coverage"]
                    + [(None if v is None else round(v / 100.0, 4))
                       for v in by_cam[cam]])
    wsd, c0, c1, nr = cd.block(cols)

    bar = BarChart()
    bar.type = "col"
    bar.title = "Daily coverage by channel, with declared outages shaded"
    bar.add_data(cd.ref(c0 + 1, c0 + 1, 1, nr), titles_from_data=True)
    bar.set_categories(cd.ref(c0, c0, 2, nr))
    _colour(bar.series[0], CLR_GREY)
    bar.dataLabels = None
    _axis(bar, "day (IST)", "share of the day", None, "0%", y_min=0, y_max=1.0)

    line = LineChart()
    line.add_data(cd.ref(c0 + 2, c1, 1, nr), titles_from_data=True)
    for s in line.series:
        s.marker = Marker(symbol="circle", size=5)
        s.smooth = False
    line.dataLabels = None
    bar += line
    _legend(bar, 1 + len(cams), position="r")
    return _size(bar, h=10.0, w=22.0)
