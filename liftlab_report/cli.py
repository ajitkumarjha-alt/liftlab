"""liftlab-report CLI.

  liftlab-report --from <ISO ts> --to <ISO ts> [--out <path>]
                 [--peak-window auto|HH:MM-HH:MM] [--population N]
                 [--min-close 0.5]
                 [--db <sqlite path>] [--gateway site-A] [--banks <json>]

Naive (offset-less) --from/--to are interpreted as IST (+05:30), the
building's clock. The DB is opened READ-ONLY (sqlite mode=ro + query_only);
this tool cannot mutate gateway tables.
"""

from __future__ import annotations

import argparse
import os
import sys
from datetime import datetime
from pathlib import Path

from . import demand_log, eras, model, reader, workbook


class ReportError(Exception):
    """Raised instead of SystemExit so the builder stays web-callable."""


def parse_ts(s: str) -> float:
    try:
        dt = datetime.fromisoformat(s)
    except ValueError as e:
        raise ReportError(f"bad timestamp {s!r}: {e}") from e
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=eras.IST)
    return dt.timestamp()


def default_out_dir() -> Path:
    p = Path("/mnt/user-data/outputs")
    if p.is_dir():
        return p
    q = Path("./reports")
    q.mkdir(parents=True, exist_ok=True)
    return q


def build_context(db_path: str, gw: str, t0: float, t1: float,
                  peak_window: str = "auto", population: int | None = None,
                  banks_path: str | None = None,
                  min_close_s: float | None = None,
                  capacity_path: str | None = None) -> dict:
    """Everything the workbook needs, precomputed. Read-only throughout."""
    if t1 <= t0:
        raise ReportError("--to must be after --from")
    db = reader.open_ro(db_path)
    try:
        cams = [c for c in reader.fetch_cams(db, gw)]
        # Install the building's own lift names before ANYTHING renders a
        # label. Unconditional, so a run never inherits the previous run's.
        channel_labels = reader.fetch_channel_labels(db, gw)
        eras.set_display_labels(channel_labels)
        validation = reader.fetch_validation(db, gw)
        stamps = reader.validation_stamps(validation)
        registry = reader.fetch_registry(db, gw)
        analyzer_versions = reader.fetch_analyzer_versions(db, gw)
        pi_cycles = reader.fetch_pi_cycles(db, gw, t0, t1)
        gpu_cycles, funnels = reader.fetch_gpu_cycles(db, gw, t0, t1)
        transits = reader.fetch_transits(db, gw, t0, t1)
        floor_status = reader.fetch_floor_read_status(db, gw, t0, t1)
        cov_buckets = reader.coverage_buckets(db, gw, cams, t0, t1)
        # DEMAND LOG source data. floor_confidence uses the door engine's own definition of a
        # confident read so the "why there is no per-floor table" line is measured, not asserted.
        floor_conf = reader.fetch_floor_confidence(db, gw, t0, t1,
                                                   demand_log.FLOOR_OK_REASONS)
        # TIER-2 evidence: per camera PER ERA, plus the confident reads the speed walk needs.
        tier2_evidence = reader.fetch_tier2_evidence(db, gw, t0, t1)
        confident_reads = reader.fetch_confident_reads(db, gw, t0, t1)
        # PEAK CAR OCCUPANCY: measured per door-open episode. Read unfiltered — the evidence rule
        # is applied in model.measured_occupancy so the workbook can report how many episodes
        # carried no measurement at all.
        occ_episodes = reader.fetch_occupancy_episodes(db, gw, t0, t1)
        occ_feature = reader.occupancy_feature_present(db)
    finally:
        db.close()

    floor_s = (eras.MIN_PLAUSIBLE_CLOSE_S if min_close_s is None
               else float(min_close_s))
    all_cycles = pi_cycles + gpu_cycles
    aggs = model.aggregate_cycles(all_cycles, t0, t1, min_close_s=floor_s)
    suspected_gaps = model.detect_suspected_gaps(all_cycles, transits, cams, t0, t1)
    transit_aggs = model.aggregate_transits(transits, stamps, t0, t1)
    dlog = demand_log.build(transits, cov_buckets, stamps, validation, cams, t0, t1,
                            reader.BUCKET_S)
    dlog_floor_note = demand_log.floor_attribution_note(floor_conf)
    floor_speed = model.floor_speed_segments(confident_reads)
    # PER-FLOOR DEMAND: join transits to the stop they happened in. Cycles already carry a floor,
    # so this is the join, not a re-walk of the door stream.
    per_floor = model.per_floor_demand(gpu_cycles, transits, stamps, tier2_evidence, t0, t1)
    # OCCUPANCY: derived, reset at every idle gap, never presented as a measurement.
    occ = model.occupancy_periods(transits, stamps, validation, t0, t1)
    occ_summary = model.occupancy_summary(occ)
    # MEASURED occupancy — a different instrument from the derived one above, per era.
    occ_measured = model.measured_occupancy(occ_episodes, stamps, t0, t1)
    capacity = eras.load_capacity(capacity_path)
    capacity_st = eras.capacity_status(capacity)
    # The evidence table behind "per-floor is not viable yet". Shipped; the per-floor numbers are not.
    attribution = model.cycle_attribution(gpu_cycles, transits, cams, t0, t1)
    coeff_blockers = model.coefficient_blockers(tier2_evidence, floor_speed, cams)
    profile = model.hourly_profile(transits, pi_cycles)
    fixed = model.parse_peak_window(peak_window)
    peaks = model.peak_by_day(transits, pi_cycles, t0, t1, fixed)
    boundaries = model.boundaries_crossed(t0, t1, aggs, transit_aggs)

    # coverage % per cam: buckets with data / non-gap buckets
    n_buckets = max(1, int((t1 - t0) // reader.BUCKET_S))
    coverage_pct, coverage_daily = {}, {"days": [], "by_cam": {}}
    from datetime import timedelta
    day0 = datetime.fromtimestamp(t0, eras.IST).replace(hour=0, minute=0,
                                                        second=0, microsecond=0)
    days = []
    d = day0
    while d.timestamp() < t1:
        days.append(d)
        d += timedelta(days=1)
    coverage_daily["days"] = [d.date().isoformat() for d in days]
    for cam in cams:
        gap_s = eras.gap_overlap_s(cam, t0, t1)
        gap_buckets = int(gap_s // reader.BUCKET_S)
        usable = max(1, n_buckets - gap_buckets)
        have = len([b for b in cov_buckets.get(cam, ())
                    if not eras.in_gap(cam, t0 + (b + 0.5) * reader.BUCKET_S)])
        coverage_pct[cam] = min(100.0, 100.0 * have / usable)
        daily = []
        for d in days:
            d0 = max(d.timestamp(), t0)
            d1 = min((d + timedelta(days=1)).timestamp(), t1)
            nb = max(1, int((d1 - d0) // reader.BUCKET_S))
            gb = int(eras.gap_overlap_s(cam, d0, d1) // reader.BUCKET_S)
            usable_d = max(0, nb - gb)
            have_d = len([b for b in cov_buckets.get(cam, ())
                          if d0 <= t0 + b * reader.BUCKET_S < d1
                          and not eras.in_gap(cam, t0 + (b + 0.5) * reader.BUCKET_S)])
            daily.append(round(100.0 * have_d / usable_d, 1) if usable_d else None)
        coverage_daily["by_cam"][cam] = daily

    # counting-version cross-check: declared epoch vs live tables
    version_mismatches = []
    cv_attribution = {}
    now_ver = {cam: eras.counting_version_at(
        cam, min(t1, datetime.now(eras.IST).timestamp()), stamps) for cam in cams}
    for cam in cams:
        src = ("camera_validation stamp" if stamps.get(cam) else "declared epochs")
        cv_attribution[cam] = src
        live = analyzer_versions.get(cam)
        if live and live != now_ver[cam]:
            version_mismatches.append(
                f"{cam}: declared/stamped attribution says {now_ver[cam]!r} but "
                f"analyzer_status reports {live!r} — update COUNTING_EPOCHS "
                f"in eras.py before trusting counting-era splits for this camera")

    banks, bank_evidence = eras.load_banks_with_derivation(
        registry, cams, banks_path)

    demand = model.demand_by_lift_hour(
        transits, pi_cycles, cov_buckets, reader.BUCKET_S, stamps,
        cams, t0, t1)

    # Fleet-wide sampling resolution, for the plain-English sheet: the eras
    # agree in practice, but say so from the data rather than assume it.
    era_quanta = {k: a["quantum_s"] for k, a in aggs.items()
                  if a["quantum_s"] is not None}
    fleet_quantum = None
    if era_quanta:
        vals = sorted(era_quanta.values())
        fleet_quantum = vals[len(vals) // 2]

    # raw rows
    raw_rows = []
    for c in all_cycles:
        ver = ("pi-watch (on-device counter)" if c["instrument"] == eras.PI_WATCH
               else eras.counting_version_at(c["cam"], c["ts"], stamps))
        raw_rows.append({
            "ts_ist": datetime.fromtimestamp(c["ts"], eras.IST).isoformat(),
            "ts_epoch": round(c["ts"], 3), "channel": c["cam"],
            "lift_label": eras.lift_label(c["cam"]),
            "bank": banks.get(c["cam"]) or "UNKNOWN",
            "instrument": c["instrument"], "counting_version": ver,
            "era_id": c["era_id"], "event_type": "door_cycle",
            "door_open_ts": (c.get("open_start_iso")
                             or datetime.fromtimestamp(c["ts"], eras.IST).isoformat()),
            "door_close_ts": (c.get("close_full_iso")
                              or (datetime.fromtimestamp(c["close_ts"], eras.IST).isoformat()
                                  if c.get("close_ts") else None)),
            "close_travel_s": c.get("close_travel_s"),
            "below_floor": ("YES" if (c.get("close_travel_s") is not None
                                      and float(c["close_travel_s"]) < floor_s)
                            else ""),
            "cycle_class": c.get("cycle_class"),
            "boarded": c.get("boarded"), "alighted": c.get("alighted"),
            "floor": c.get("floor"), "floor_source": c.get("floor_source"),
            "precision_at_time": reader.precision_str(validation, c["cam"], ver),
            "in_gap": "YES" if eras.in_gap(c["cam"], c["ts"]) else "",
        })
    for t in transits:
        ver = eras.counting_version_at(t["cam"], t["ts"], stamps)
        raw_rows.append({
            "ts_ist": datetime.fromtimestamp(t["ts"], eras.IST).isoformat(),
            "ts_epoch": round(t["ts"], 3), "channel": t["cam"],
            "lift_label": eras.lift_label(t["cam"]),
            "bank": banks.get(t["cam"]) or "UNKNOWN",
            "instrument": t["instrument"], "counting_version": ver,
            "era_id": f"counting:{ver}", "event_type": "transit",
            "door_open_ts": None, "door_close_ts": None, "close_travel_s": None,
            "below_floor": "", "cycle_class": None,
            "boarded": 1 if t["direction"] == "in" else None,
            "alighted": 1 if t["direction"] == "out" else None,
            "floor": None, "floor_source": None,
            "precision_at_time": reader.precision_str(validation, t["cam"], ver),
            "in_gap": "YES" if eras.in_gap(t["cam"], t["ts"]) else "",
        })
    raw_rows.sort(key=lambda r: r["ts_epoch"])

    ctx = {
        "t0": t0, "t1": t1,
        "from_iso": datetime.fromtimestamp(t0, eras.IST).isoformat(),
        "to_iso": datetime.fromtimestamp(t1, eras.IST).isoformat(),
        "generated_at": datetime.now(eras.IST).isoformat(timespec="seconds"),
        "gw": gw, "db_path": str(db_path), "cams": cams, "banks": banks,
        "validation": validation, "registry": registry,
        "demand_log": dlog, "demand_log_floor_note": dlog_floor_note,
        "floor_confidence": floor_conf,
        "tier2_evidence": tier2_evidence, "floor_speed": floor_speed,
        "per_floor_demand": per_floor, "occupancy": occ, "occupancy_summary": occ_summary,
        "measured_occupancy": occ_measured, "occupancy_episodes": occ_episodes,
        "occupancy_feature_present": occ_feature,
        "capacity": capacity, "capacity_status": capacity_st,
        "cycle_attribution": attribution,
        "coefficient_blockers": coeff_blockers,
        "analyzer_versions": analyzer_versions,
        "aggs": aggs, "transit_aggs": transit_aggs, "funnels": funnels,
        "boundaries": boundaries, "profile": profile, "peaks": peaks,
        "floor_status": floor_status, "coverage_pct": coverage_pct,
        "coverage_daily": coverage_daily, "population": population,
        "peak_window_spec": peak_window, "all_cycles": all_cycles,
        "raw_rows": raw_rows, "n_raw_rows": len(raw_rows),
        "version_mismatches": version_mismatches,
        "cv_attribution": cv_attribution,
        "stamps": stamps,
        "min_close_s": floor_s,
        "canonical": model.canonical_figures(demand, peaks, aggs, coverage_pct),
        "channel_labels": channel_labels,
        "unlabelled_cams": [c for c in cams if not eras.has_display_label(c)],
        "bank_evidence": bank_evidence,
        "demand": demand,
        "suspected_gaps": suspected_gaps,
        "era_quanta": era_quanta,
        "fleet_quantum_s": fleet_quantum,
        "transits": transits,
    }
    # Built last: the provenance lines quote ctx's own range/gateway/precision fields, and
    # they must be identical in the CSV and the sheet so the two can never disagree.
    ctx["demand_log_notes"] = demand_log.notes(ctx, dlog) + [dlog_floor_note]
    return ctx


def build_workbook(ctx):
    """Data sheets first — they record WHERE they put each figure in `anchors`.
    READ THIS FIRST is built last from those anchors, so every 'Evidence:' line
    on it names a cell range that really holds the number it quotes, then it is
    moved to the front where a first-time reader will meet it."""
    from openpyxl import Workbook
    wb = Workbook()
    cd = workbook.ChartData(wb)
    anchors: dict = {}
    workbook.sheet_summary(wb, ctx, cd, anchors)
    workbook.sheet_vs_sheet(wb, ctx, cd, anchors)
    workbook.sheet_per_lift(wb, ctx, cd, anchors)
    workbook.sheet_fleet(wb, ctx, cd, anchors)
    workbook.sheet_demand(wb, ctx, cd, anchors)
    workbook.sheet_demand_log(wb, ctx, anchors)
    workbook.sheet_peak(wb, ctx, cd, anchors)
    workbook.sheet_raw(wb, ctx, anchors)
    workbook.sheet_coverage(wb, ctx, cd, anchors)
    workbook.sheet_car_loading(wb, ctx, anchors)
    workbook.sheet_attribution(wb, ctx, anchors)
    workbook.sheet_tier2(wb, ctx, anchors)
    workbook.sheet_read_this_first(wb, ctx, anchors)
    wb.move_sheet("READ THIS FIRST", offset=-(len(wb.sheetnames) - 1))
    # keep DATA_CHARTS last in the tab order
    wb.move_sheet("DATA_CHARTS", offset=len(wb.sheetnames))
    return wb


def run_all(argv=None) -> list:
    ap = argparse.ArgumentParser(
        prog="liftlab-report",
        description="Export an era-hygienic Excel workbook of lift analytics "
                    "from the gateway DB (read-only).")
    ap.add_argument("--from", dest="from_ts", required=True,
                    help="range start, ISO timestamp (naive = IST)")
    ap.add_argument("--to", dest="to_ts", required=True,
                    help="range end (exclusive), ISO timestamp (naive = IST)")
    ap.add_argument("--out", default=None, help="output .xlsx path")
    ap.add_argument("--format", dest="fmt", choices=("csv", "xlsx", "both"),
                    default="both",
                    help="which outputs to write. 'csv' writes the DEMAND LOG as two CSVs "
                         "(hourly + daily rollup) and skips the workbook; 'xlsx' writes only "
                         "the workbook (which contains the same rows as a DEMAND LOG sheet); "
                         "'both' (default) writes both.")
    ap.add_argument("--peak-window", default="auto",
                    help="'auto' (worst 5-min window per day) or fixed HH:MM-HH:MM")
    ap.add_argument("--population", type=int, default=None,
                    help="tower population for peak-demand %% vs the 8%% HC assumption")
    ap.add_argument("--min-close", dest="min_close", type=float,
                    default=eras.MIN_PLAUSIBLE_CLOSE_S,
                    help="one-frame quantization floor in seconds: closes below "
                         "this are door-state flip artifacts and are rejected "
                         f"from every close-travel statistic "
                         f"(default {eras.MIN_PLAUSIBLE_CLOSE_S})")
    ap.add_argument("--db", default=os.environ.get("GATEWAY_DB",
                                                   "/var/lib/liftlab/gateway.db"),
                    help="gateway sqlite DB path (default: $GATEWAY_DB)")
    ap.add_argument("--gateway", default=os.environ.get("GATEWAY_ID", "site-A"))
    ap.add_argument("--banks", default=None,
                    help="bank sidecar json (default: lift_banks.json beside the module)")
    ap.add_argument("--capacity", default=None,
                    help=f"car capacity sidecar json (default: {eras.CAPACITY_SIDECAR_DEFAULT} "
                         f"beside the module). Must state persons per car AND the BASIS those "
                         f"persons are counted on; without a confirmed basis the workbook prints "
                         f"measured occupancy in absolute people and withholds the loading factor")
    args = ap.parse_args(argv)

    t0, t1 = parse_ts(args.from_ts), parse_ts(args.to_ts)
    if not Path(args.db).exists():
        raise ReportError(f"DB not found: {args.db}")
    if args.min_close < 0:
        raise ReportError("--min-close must be >= 0")
    ctx = build_context(args.db, args.gateway, t0, t1,
                        peak_window=args.peak_window,
                        population=args.population, banks_path=args.banks,
                        min_close_s=args.min_close, capacity_path=args.capacity)
    if args.out:
        out = Path(args.out)
        out.parent.mkdir(parents=True, exist_ok=True)
    else:
        stamp = datetime.now(eras.IST).strftime("%Y%m%d-%H%M%S")
        f0 = datetime.fromtimestamp(t0, eras.IST).strftime("%Y%m%d")
        f1 = datetime.fromtimestamp(t1, eras.IST).strftime("%Y%m%d")
        out = default_out_dir() / f"liftlab-report_{f0}-{f1}_{stamp}.xlsx"

    written = []
    if args.fmt in ("csv", "both"):
        written += demand_log.write_csvs(out, ctx, ctx["demand_log"],
                                         ctx["demand_log_floor_note"])
    if args.fmt in ("xlsx", "both"):
        wb = build_workbook(ctx)
        wb.save(out)
        written.append(out)
    return written


def run(argv=None) -> Path:
    """The PRIMARY artifact, as a single Path.

    Kept returning one Path because prove_report.py and any web caller depend on that contract;
    run_all() is the one that reports everything written. When a workbook was produced it is the
    primary (it embeds the same DEMAND LOG rows as a sheet); under --format csv the hourly CSV is."""
    written = run_all(argv)
    xlsx = [p for p in written if p.suffix == ".xlsx"]
    return xlsx[0] if xlsx else written[0]


def main(argv=None) -> int:
    try:
        written = run_all(argv)
    except ReportError as e:
        print(f"liftlab-report: {e}", file=sys.stderr)
        return 2
    for p in written:
        print(f"wrote {p}")
    return 0
